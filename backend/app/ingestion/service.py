import hashlib
import json
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pymupdf
import yaml
from openpyxl import load_workbook
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.service import AuthService
from app.calculations import classify_ticket_severity, evaluate_ticket_sla
from app.config import Settings
from app.db.models import (
    Account,
    DatasetSnapshot,
    Document,
    DocumentChunk,
    IngestionRun,
    Order,
    Ticket,
    User,
    UserAccountScope,
)
from app.services.embeddings import LocalEmbeddingService


AUTHORITY_RANKS = {
    "customer_agreement": 100,
    "current_policy": 90,
    "current_sop": 85,
    "product_guide": 75,
    "deprecated_policy": 10,
}


class IngestionError(RuntimeError):
    pass


class IngestionService:
    def __init__(self, settings: Settings):
        self.settings = settings
        self._embedding_model: LocalEmbeddingService | None = None

    @staticmethod
    def json_safe(value: Any) -> Any:
        return json.loads(json.dumps(value, default=str))

    @staticmethod
    def sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest()

    @staticmethod
    def parse_source_datetime(value: Any, timezone: ZoneInfo) -> datetime | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            parsed = value
        else:
            parsed = datetime.strptime(str(value).strip(), "%Y-%m-%d %H:%M")
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone)
        return parsed.astimezone(UTC)

    @staticmethod
    def populated_rows(worksheet: Any) -> list[list[Any]]:
        rows: list[list[Any]] = []
        for row in worksheet.iter_rows(values_only=True):
            values = list(row)
            if any(value is not None for value in values):
                rows.append(values)
        return rows

    @staticmethod
    def rows_as_dicts(rows: list[list[Any]]) -> list[dict[str, Any]]:
        if not rows:
            return []
        headers = [str(value).strip() for value in rows[0]]
        return [dict(zip(headers, row, strict=False)) for row in rows[1:]]

    def load_manifest(self) -> dict[str, Any]:
        if not self.settings.data_manifest_path.exists():
            raise IngestionError(f"Manifest not found: {self.settings.data_manifest_path}")
        with self.settings.data_manifest_path.open(encoding="utf-8") as stream:
            manifest = yaml.safe_load(stream)
        if not isinstance(manifest, dict) or not isinstance(manifest.get("documents"), list):
            raise IngestionError("Manifest must contain a documents list")
        return manifest

    def validate_files(self, manifest: dict[str, Any]) -> dict[str, str]:
        required = [entry["filename"] for entry in manifest["documents"]]
        required.append("ParcelPilot_Assessment_Data.xlsx")
        checksums: dict[str, str] = {}
        for filename in required:
            path = self.settings.data_pack_dir / filename
            if not path.exists() or not path.is_file() or path.stat().st_size == 0:
                raise IngestionError(f"Required source file is missing or empty: {filename}")
            checksums[filename] = self.sha256(path)
        return checksums

    def load_embedding_model(self) -> LocalEmbeddingService:
        if self._embedding_model is None:
            self._embedding_model = LocalEmbeddingService(self.settings.embedding_model)
        return self._embedding_model

    @staticmethod
    def clean_pdf_text(text: str) -> str:
        text = text.replace("\u00a0", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def chunk_page(self, text: str) -> list[tuple[str | None, str]]:
        target_chars = self.settings.rag_chunk_tokens * 4
        overlap_chars = self.settings.rag_chunk_overlap_tokens * 4
        if len(text) <= target_chars:
            heading_match = re.search(r"(?m)^\d+\.\s+[^\n]+", text)
            return [(heading_match.group(0) if heading_match else None, text)]
        chunks: list[tuple[str | None, str]] = []
        start = 0
        while start < len(text):
            end = min(len(text), start + target_chars)
            if end < len(text):
                boundary = text.rfind("\n", start, end)
                if boundary > start + target_chars // 2:
                    end = boundary
            content = text[start:end].strip()
            heading_match = re.search(r"(?m)^\d+\.\s+[^\n]+", content)
            if content:
                chunks.append((heading_match.group(0) if heading_match else None, content))
            if end >= len(text):
                break
            start = max(end - overlap_chars, start + 1)
        return chunks

    async def reset_imported_data(self, db: AsyncSession) -> None:
        await db.execute(delete(Document))
        await db.execute(delete(UserAccountScope))
        await db.execute(delete(User))
        await db.execute(delete(Account))
        await db.execute(delete(DatasetSnapshot))
        await db.commit()

    async def ingest(self, db: AsyncSession, *, reset: bool = False) -> dict[str, Any]:
        manifest = self.load_manifest()
        checksums = self.validate_files(manifest)
        run = IngestionRun(
            status="running",
            input_manifest=self.json_safe(manifest),
            source_checksums=checksums,
            embedding_model=self.settings.embedding_model,
            parser_versions={"pymupdf": pymupdf.VersionBind, "openpyxl": "3.1"},
        )
        db.add(run)
        await db.commit()
        await db.refresh(run)
        run_id = run.id

        try:
            if reset:
                await self.reset_imported_data(db)

            workbook_path = self.settings.data_pack_dir / "ParcelPilot_Assessment_Data.xlsx"
            workbook = load_workbook(workbook_path, data_only=True, read_only=False)
            required_sheets = {"README", "accounts", "orders", "tickets"}
            missing = required_sheets.difference(workbook.sheetnames)
            if missing:
                raise IngestionError(f"Workbook is missing sheets: {sorted(missing)}")

            readme = self.populated_rows(workbook["README"])
            readme_values = {str(row[0]).strip(): row[1] for row in readme[1:] if len(row) >= 2}
            snapshot_raw = readme_values.get("Dataset snapshot")
            if not snapshot_raw:
                raise IngestionError("Workbook README is missing Dataset snapshot")
            snapshot_match = re.fullmatch(r"(.+)\s+(Asia/Kolkata)", str(snapshot_raw).strip())
            if not snapshot_match:
                raise IngestionError("Dataset snapshot must include Asia/Kolkata timezone")
            source_timezone = ZoneInfo(snapshot_match.group(2))
            snapshot_at = self.parse_source_datetime(snapshot_match.group(1), source_timezone)
            if snapshot_at is None:
                raise IngestionError("Dataset snapshot could not be parsed")

            await db.execute(update(DatasetSnapshot).values(is_active=False))
            snapshot = DatasetSnapshot(
                snapshot_at=snapshot_at,
                source_filename=workbook_path.name,
                source_sha256=checksums[workbook_path.name],
                is_active=True,
                metadata_json=self.json_safe(
                    {"currency": readme_values.get("Currency"), "notes": readme_values}
                ),
            )
            db.add(snapshot)
            await db.flush()

            account_rows = self.rows_as_dicts(self.populated_rows(workbook["accounts"]))
            accounts_by_external: dict[str, Account] = {}
            for source_row, payload in enumerate(account_rows, start=2):
                external_id = str(payload.get("account_id") or "").strip()
                if not external_id:
                    raise IngestionError(f"accounts!{source_row} is missing account_id")
                account = Account(
                    external_id=external_id,
                    name=str(payload.get("account_name") or "").strip(),
                    plan=str(payload.get("plan") or "").strip(),
                    status=str(payload.get("status") or "").strip(),
                    csm=payload.get("csm"),
                    contract_file=payload.get("contract_file"),
                    premium_support=bool(payload.get("premium_support")),
                    notes=payload.get("notes"),
                    dataset_snapshot_id=snapshot.id,
                    source_sheet="accounts",
                    source_row=source_row,
                    raw_payload=self.json_safe(payload),
                )
                db.add(account)
                accounts_by_external[external_id] = account
            await db.flush()

            order_rows = self.rows_as_dicts(self.populated_rows(workbook["orders"]))
            for source_row, payload in enumerate(order_rows, start=2):
                account_external_id = str(payload.get("account_id") or "").strip()
                account = accounts_by_external.get(account_external_id)
                if account is None:
                    raise IngestionError(f"orders!{source_row} references unknown account")
                db.add(
                    Order(
                        external_id=str(payload["order_id"]).strip(),
                        account_id=account.id,
                        carrier=str(payload.get("carrier") or "").strip(),
                        status=str(payload.get("status") or "").strip(),
                        booked_at=self.parse_source_datetime(payload.get("booked_at"), source_timezone),
                        pickup_window_start=self.parse_source_datetime(payload.get("pickup_window_start"), source_timezone),
                        pickup_window_end=self.parse_source_datetime(payload.get("pickup_window_end"), source_timezone),
                        pickup_actual_at=self.parse_source_datetime(payload.get("pickup_actual_at"), source_timezone),
                        shipment_fee_inr=payload.get("shipment_fee_inr") or 0,
                        carrier_fault=payload.get("carrier_fault"),
                        customer_fault=payload.get("customer_fault"),
                        cancellation_requested_at=self.parse_source_datetime(payload.get("cancellation_requested_at"), source_timezone),
                        notes=payload.get("notes"),
                        dataset_snapshot_id=snapshot.id,
                        source_sheet="orders",
                        source_row=source_row,
                        raw_payload=self.json_safe(payload),
                    )
                )

            ticket_rows = self.rows_as_dicts(self.populated_rows(workbook["tickets"]))
            for source_row, payload in enumerate(ticket_rows, start=2):
                account_external_id = str(payload.get("account_id") or "").strip()
                account = accounts_by_external.get(account_external_id)
                if account is None:
                    raise IngestionError(f"tickets!{source_row} references unknown account")
                created_at = self.parse_source_datetime(payload.get("created_at"), source_timezone)
                severity = classify_ticket_severity(
                    str(payload.get("subject") or ""), str(payload.get("description") or "")
                )
                sla = evaluate_ticket_sla(
                    account_external_id=account_external_id,
                    plan=account.plan,
                    severity=severity,
                    created_at=created_at,
                    snapshot_at=snapshot_at,
                )
                due_raw = sla.get("due_at")
                db.add(
                    Ticket(
                        external_id=str(payload["ticket_id"]).strip(),
                        account_id=account.id,
                        source_created_at=created_at,
                        status=str(payload.get("status") or "").strip(),
                        subject=str(payload.get("subject") or "").strip(),
                        description=str(payload.get("description") or "").strip(),
                        channel=str(payload.get("channel") or "").strip(),
                        assigned_to=payload.get("assigned_to"),
                        last_customer_message_at=self.parse_source_datetime(payload.get("last_customer_message_at"), source_timezone),
                        historical_resolution=payload.get("historical_resolution"),
                        severity=severity,
                        sla_due_at=datetime.fromisoformat(due_raw) if due_raw else None,
                        dataset_snapshot_id=snapshot.id,
                        source_sheet="tickets",
                        source_row=source_row,
                        raw_payload=self.json_safe({**payload, "sla_evaluation": sla}),
                    )
                )

            await db.flush()
            await self.ingest_documents(db, manifest, checksums, accounts_by_external, snapshot_at)
            await self.seed_demo_users(db, accounts_by_external)

            run.status = "completed"
            run.completed_at = datetime.now(UTC)
            run.imported_count = len(account_rows) + len(order_rows) + len(ticket_rows) + len(manifest["documents"])
            run.activated_snapshot_id = snapshot.id
            await db.commit()
            return {
                "status": "completed",
                "snapshot_at": snapshot_at.isoformat(),
                "accounts": len(account_rows),
                "orders": len(order_rows),
                "tickets": len(ticket_rows),
                "documents": len(manifest["documents"]),
            }
        except Exception as exc:
            await db.rollback()
            result = await db.execute(select(IngestionRun).where(IngestionRun.id == run_id))
            failed_run = result.scalar_one_or_none()
            if failed_run:
                failed_run.status = "failed"
                failed_run.completed_at = datetime.now(UTC)
                failed_run.failed_count = 1
                failed_run.validation_errors = [{"error": type(exc).__name__, "message": str(exc)}]
                await db.commit()
            raise

    async def ingest_documents(
        self,
        db: AsyncSession,
        manifest: dict[str, Any],
        checksums: dict[str, str],
        accounts_by_external: dict[str, Account],
        snapshot_at: datetime,
    ) -> None:
        model = self.load_embedding_model()
        for entry in manifest["documents"]:
            path = self.settings.data_pack_dir / entry["filename"]
            account = accounts_by_external.get(entry.get("account_external_id"))
            pdf = pymupdf.open(path)
            document = Document(
                filename=path.name,
                title=entry["title"],
                document_type=entry["document_type"],
                version=entry.get("version"),
                status=entry["status"],
                effective_from=entry.get("effective_from"),
                effective_to=entry.get("effective_to"),
                account_id=account.id if account else None,
                authority_class=entry["authority_class"],
                authority_rank=AUTHORITY_RANKS[entry["authority_class"]],
                sha256=checksums[path.name],
                page_count=pdf.page_count,
                metadata_json={
                    "account_external_id": entry.get("account_external_id"),
                    "dataset_snapshot_at": snapshot_at.isoformat(),
                },
            )
            db.add(document)
            await db.flush()
            chunk_index = 0
            for page_number, page in enumerate(pdf, start=1):
                text = self.clean_pdf_text(page.get_text("text", sort=True))
                if not text:
                    raise IngestionError(f"No extractable text in {path.name} page {page_number}")
                for heading, content in self.chunk_page(text):
                    vector = model.embed_one(content)
                    db.add(
                        DocumentChunk(
                            document_id=document.id,
                            chunk_index=chunk_index,
                            page_start=page_number,
                            page_end=page_number,
                            section_heading=heading,
                            content=content,
                            token_count=max(1, len(content) // 4),
                            embedding=vector,
                            metadata_json={"filename": path.name},
                        )
                    )
                    chunk_index += 1
            pdf.close()

    async def seed_demo_users(
        self, db: AsyncSession, accounts_by_external: dict[str, Account]
    ) -> None:
        auth = AuthService(self.settings)
        users = [
            ("northstar@demo.parcelpilot.com", "Northstar Customer", "customer", "ACCT-001"),
            ("lumenworks@demo.parcelpilot.com", "LumenWorks Customer", "customer", "ACCT-002"),
            ("beacon@demo.parcelpilot.com", "Beacon Customer", "customer", "ACCT-003"),
            ("support@demo.parcelpilot.com", "ParcelPilot Support", "support_agent", None),
            ("ops@demo.parcelpilot.com", "ParcelPilot Operations", "operations_manager", None),
        ]
        for email, name, role, account_external_id in users:
            user = User(
                email=email,
                display_name=name,
                role=role,
                password_hash=auth.hash_password(self.settings.demo_user_password),
            )
            db.add(user)
            await db.flush()
            if account_external_id:
                db.add(
                    UserAccountScope(
                        user_id=user.id,
                        account_id=accounts_by_external[account_external_id].id,
                        scope="read",
                    )
                )
